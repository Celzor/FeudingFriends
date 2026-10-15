#!/usr/bin/env python3
"""Import a supplied workbook and the question section of a ClassPoint PDF.

Requires openpyxl and pdftotext (Poppler). External files are data, never commands.
No Fast Money sheet is imported. No network access or spreadsheet write occurs.
"""
import argparse
import collections
import csv
import hashlib
import json
import re
import subprocess
import tempfile
from pathlib import Path
from openpyxl import load_workbook
from clean_questions import ROOT, clean, edit, key, write


def value(cell):
    if cell is None:
        return ''
    if isinstance(cell, float):
        return format(cell, '.15g')
    return str(cell)


def parse_pdf(path):
    text = subprocess.run(['pdftotext', '-layout', str(path), '-'], check=True, capture_output=True, text=True).stdout
    text = text.replace('\f', '\n').split('Turn these Family Feud questions')[0]
    headings = list(re.finditer(r'^\s*(\d+)\.\s*', text, re.M))
    records = []
    for index, heading in enumerate(headings):
        block = text[heading.end(): headings[index+1].start() if index+1 < len(headings) else len(text)]
        question, answers = [], []
        for line in [line.strip() for line in block.splitlines() if line.strip()]:
            answer = re.fullmatch(r'(.*?)\s*\((\d+)\)', line)
            if answer:
                answers.append({'text': answer[1].strip(), 'score': int(answer[2])})
            elif not answers:
                question.append(line)
            else:
                raise ValueError(f'Unparsed text in PDF question {heading[1]}: {line!r}')
        if not 3 <= len(answers) <= 8:
            raise ValueError(f'Invalid PDF answer count for question {heading[1]}')
        records.append({'id': int(heading[1]), 'question': ' '.join(question), 'answers': answers})
    if [r['id'] for r in records] != list(range(1,132)):
        raise ValueError('Expected all 131 sequential ClassPoint questions; check the PDF layout before importing')
    return records


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook',type=Path,required=True)
    parser.add_argument('--pdf',type=Path,required=True)
    parser.add_argument('--output-dir',type=Path,default=ROOT/'server/data')
    args=parser.parse_args()
    workbook=load_workbook(args.workbook,read_only=True,data_only=True)
    raw=collections.defaultdict(list)
    sheets={}; excluded={}
    for sheet in workbook:
        records=[list(row) for row in sheet.iter_rows(values_only=True) if any(v is not None for v in row)]
        if 'fast money' in sheet.title.casefold():
            excluded[sheet.title]=len(records)-1
            continue
        if sheet.title not in ['3 Answers','4 Answers','5 Answers','6 Answers','7 Answers','Sheet1','No Points 3','No Points 4','No Points 5','No Points 6','No Points 7']:
            raise ValueError(f'Unreviewed worksheet: {sheet.title}')
        sheets[sheet.title]=len(records)-1
        for line,row in enumerate(records[1:],2):
            while row and row[-1] is None:
                row.pop()
            if not row or row[0] is None:
                continue
            values=list(map(value,row))
            if sheet.title.startswith('No Points '):
                count=int(sheet.title.split()[-1]);pool={4:1,5:2}.get(count,0)
                values=values[:count+1]
                if pool==0:
                    values+=['']*(8-len(values))
            else:
                if len(values)%2!=1:
                    raise ValueError(f'{sheet.title}:{line}: incomplete answer/score pair')
                pool=(len(values)-1)//2
                if not 3<=pool<=8:
                    raise ValueError(f'{sheet.title}:{line}: invalid answer count')
            raw[pool].append(values)
    pdf=parse_pdf(args.pdf)
    for record in pdf:
        row=[record['question']]
        for answer in record['answers']:
            row.extend([answer['text'],str(answer['score'])])
        raw[len(record['answers'])].append(row)
    with tempfile.TemporaryDirectory(prefix='feud-source-import-') as directory:
        staging=Path(directory)
        for pool in range(9):
            count=pool if pool>=3 else (7 if pool==0 else pool+3)
            header=['Question']
            for i in range(1,count+1):
                header.append(f'Answer {i}')
                if pool>=3:header.append(f'#{i}')
            with (staging/f'{pool}.csv').open('w',newline='') as stream:
                writer=csv.writer(stream,lineterminator='\n');writer.writerow(header);writer.writerows(raw[pool])
        boards,stats,removed=clean(staging)
    write(boards,args.output_dir)
    canonical={key(board['question']):board for board in boards}
    represented=[]
    for record in pdf:
        question=edit(record['question'],question=True)
        board=canonical.get(key(question))
        if board is None:raise ValueError(f'PDF question {record["id"]} missing after reconciliation: {question}')
        represented.append({'number':record['id'],'question':board['question'],'pool':board['pool']})
    report={
        'review_date':'2026-10-15',
        'sources':{
            'workbook':{'name':args.workbook.name,'sha256':hashlib.sha256(args.workbook.read_bytes()).hexdigest(),'exported_at':None,'worksheets':sheets},
            'classpoint_pdf':{'name':args.pdf.name,'sha256':hashlib.sha256(args.pdf.read_bytes()).hexdigest(),'numbered_questions':len(pdf)},
        },
        'excluded_fast_money_worksheets':excluded,'counts':stats,
        'pool_counts':dict(sorted(collections.Counter(board['pool'] for board in boards).items())),
        'excluded':removed,
    }
    (args.output_dir/'quality-report.json').write_text(json.dumps(report,indent=2,ensure_ascii=False)+'\n')
    (args.output_dir/'classpoint-import.json').write_text(json.dumps(represented,indent=2,ensure_ascii=False)+'\n')
    print(json.dumps(report,indent=2,ensure_ascii=False))


if __name__=='__main__':main()
